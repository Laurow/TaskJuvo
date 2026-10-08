"""Create an annotated PowerPoint walkthrough for the TaskJuvo prototype."""
from __future__ import annotations

import math
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
ARTIFACTS = ROOT / "artifacts"
ASSETS = Path("/tmp/taskjuvo-ppt-assets")
ASSETS.mkdir(parents=True, exist_ok=True)

SLIDE_W = 13.333
SLIDE_H = 7.5
NAVY = RGBColor(23, 36, 59)
MUTED = RGBColor(93, 106, 126)
ORANGE = RGBColor(184, 67, 37)
RED = RGBColor(214, 48, 37)
RED_LIGHT = RGBColor(253, 238, 234)
CANVAS = RGBColor(247, 248, 250)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(221, 226, 233)
GREEN = RGBColor(42, 104, 78)


def crop_source(name: str, bottom: int) -> Path:
    source = ARTIFACTS / name
    output = ASSETS / name
    with Image.open(source) as image:
        width, height = image.size
        image.crop((0, 0, width, min(bottom, height))).save(output, quality=94)
    return output


SCREENS = {
    "home": crop_source("TaskJuvo-homepage.png", 940),
    "wizard": crop_source("TaskJuvo-task-wizard.png", 1233),
    "matches": crop_source("TaskJuvo-matches.png", 1270),
    "profile": crop_source("TaskJuvo-profile.png", 1140),
    "confirmation": crop_source("TaskJuvo-confirmation.png", 1300),
    "workspace": crop_source("TaskJuvo-workspace.png", 1320),
    "feedback": crop_source("TaskJuvo-feedback.png", 1264),
}


def add_text(slide, x, y, w, h, text, size=14, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = Inches(0.02)
    box.text_frame.margin_right = Inches(0.02)
    box.text_frame.margin_top = Inches(0.01)
    box.text_frame.margin_bottom = Inches(0.01)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def rounded_box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True, width=1.0):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    return shape


def add_header(slide, step, title, subtitle):
    add_text(slide, 0.62, 0.28, 1.6, 0.28, f"TASKJUVO · WALKTHROUGH {step:02d}", 9.5, ORANGE, True)
    add_text(slide, 0.62, 0.58, 11.8, 0.5, title, 26, NAVY, True)
    add_text(slide, 0.62, 1.10, 11.7, 0.38, subtitle, 12.5, MUTED)


def add_footer(slide, active):
    labels = ["1 Brief", "2 Define", "3 Compare", "4 Confirm", "5 Work", "6 Reflect"]
    add_text(slide, 0.62, 7.12, 1.25, 0.18, "NAVIGATION", 7.5, MUTED, True)
    x = 1.82
    for index, label in enumerate(labels, start=1):
        is_active = index == active
        width = 1.42 if index < 6 else 1.15
        rounded_box(slide, x, 7.045, width, 0.32, RED_LIGHT if is_active else CANVAS,
                    ORANGE if is_active else LINE, width=0.8)
        add_text(slide, x + 0.06, 7.065, width - 0.12, 0.18, label, 8.7,
                 ORANGE if is_active else MUTED, is_active, PP_ALIGN.CENTER)
        x += width + 0.09
    add_text(slide, 11.7, 7.12, 1.0, 0.18, "TaskJuvo", 8, MUTED, False, PP_ALIGN.RIGHT)


def add_slide_base(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = CANVAS
    return slide


def add_screen(slide, path, x, y, w, max_h, border=LINE):
    with Image.open(path) as image:
        iw, ih = image.size
    ratio = ih / iw
    h = min(w * ratio, max_h)
    actual_w = h / ratio
    x = x + (w - actual_w) / 2
    rounded_box(slide, x - 0.045, y - 0.045, actual_w + 0.09, h + 0.09, WHITE, border, True, 0.8)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(actual_w), height=Inches(h))
    return {"x": x, "y": y, "w": actual_w, "h": h, "iw": iw, "ih": ih}


def screen_point(screen, px, py):
    return screen["x"] + (px / screen["iw"]) * screen["w"], screen["y"] + (py / screen["ih"]) * screen["h"]


def screen_box(slide, screen, left, top, right, bottom, color=RED, dash=False):
    x1, y1 = screen_point(screen, left, top)
    x2, y2 = screen_point(screen, right, bottom)
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                   Inches(x1), Inches(y1), Inches(x2 - x1), Inches(y2 - y1))
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(2.25)
    if dash:
        shape.line.dash_style = 1
    return shape


def arrow(slide, x1, y1, x2, y2, color=RED, width=2.25):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 0.12
    for turn in (math.pi - 0.48, math.pi + 0.48):
        hx = x2 + length * math.cos(angle + turn)
        hy = y2 + length * math.sin(angle + turn)
        head = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x2), Inches(y2), Inches(hx), Inches(hy))
        head.line.color.rgb = color
        head.line.width = Pt(width)
    return line


def callout(slide, x, y, w, text, number=None):
    rounded_box(slide, x, y, w, 0.38, WHITE, RED, True, 1.2)
    if number is not None:
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.08), Inches(y + 0.075), Inches(0.23), Inches(0.23))
        badge.fill.solid(); badge.fill.fore_color.rgb = RED
        badge.line.color.rgb = RED
        add_text(slide, x + 0.08, y + 0.082, 0.23, 0.18, str(number), 8.5, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, x + 0.38, y + 0.075, w - 0.46, 0.23, text, 9.5, NAVY, True)
    else:
        add_text(slide, x + 0.14, y + 0.075, w - 0.22, 0.23, text, 9.5, NAVY, True)


def add_cover(prs):
    slide = add_slide_base(prs)
    add_text(slide, 0.72, 0.62, 3.8, 0.3, "TASKJUVO · PRODUCT WALKTHROUGH", 10, ORANGE, True)
    add_text(slide, 0.72, 1.35, 8.9, 1.18, "From a business need\nto a shared outcome.", 34, NAVY, True)
    add_text(slide, 0.75, 2.73, 6.5, 0.48, "A visual guide to the clickable prototype and its main navigation path.", 15, MUTED)
    rounded_box(slide, 0.75, 3.55, 11.75, 1.42, WHITE, LINE, True, 1)
    steps = [("01", "Brief", "Start with the task"), ("02", "Define", "Make the outcome clear"),
             ("03", "Compare", "Review evidence"), ("04", "Confirm", "Agree the plan"),
             ("05", "Work", "Review progress"), ("06", "Reflect", "Improve the experience")]
    x = 1.03
    for number, label, description in steps:
        add_text(slide, x, 3.86, 0.5, 0.26, number, 10, ORANGE, True)
        add_text(slide, x, 4.17, 1.58, 0.25, label, 15, NAVY, True)
        add_text(slide, x, 4.46, 1.58, 0.29, description, 9.5, MUTED)
        if number != "06":
            arrow(slide, x + 1.62, 4.22, x + 1.93, 4.22, ORANGE, 1.55)
        x += 1.92
    add_text(slide, 0.76, 6.58, 5.9, 0.24, "Red outlines and arrows show the recommended navigation sequence.", 11, MUTED)
    add_text(slide, 11.1, 6.58, 1.35, 0.24, "6-step tour", 11, ORANGE, True, PP_ALIGN.RIGHT)
    return slide


def slide_start(prs):
    slide = add_slide_base(prs)
    add_header(slide, 1, "Start with a clear task", "The homepage gives the business owner two obvious ways into the flow.")
    home = add_screen(slide, SCREENS["home"], 0.62, 1.68, 5.72, 4.9)
    wizard = add_screen(slide, SCREENS["wizard"], 6.98, 1.68, 5.72, 4.9)
    screen_box(slide, home, 1215, 15, 1360, 70)
    callout(slide, 0.85, 6.66, 2.45, "Click the primary CTA", 1)
    arrow(slide, 3.25, 6.85, *screen_point(home, 1290, 45), RED, 1.8)
    screen_box(slide, wizard, 1230, 138, 1400, 205)
    callout(slide, 7.2, 6.66, 2.72, "Use the sample or brief", 2)
    arrow(slide, 9.93, 6.85, *screen_point(wizard, 1315, 170), RED, 1.8)
    arrow(slide, 6.42, 4.05, 6.83, 4.05, RED, 2.4)
    add_text(slide, 5.9, 3.54, 1.1, 0.25, "NEXT", 8.5, RED, True, PP_ALIGN.CENTER)
    add_footer(slide, 1)
    return slide


def slide_define(prs):
    slide = add_slide_base(prs)
    add_header(slide, 2, "Define the brief before matching", "The wizard turns a vague request into a focused, reviewable task.")
    wizard = add_screen(slide, SCREENS["wizard"], 0.62, 1.55, 8.15, 5.33)
    rounded_box(slide, 9.08, 1.7, 3.55, 4.95, WHITE, LINE, True, 1)
    add_text(slide, 9.42, 2.03, 2.85, 0.28, "Recommended path", 15, NAVY, True)
    add_text(slide, 9.42, 2.42, 2.8, 0.75, "Complete the five steps, then review the summary before moving to matches.", 12, MUTED)
    items = [("1", "Task", "Name the challenge"), ("2", "Outcome", "List deliverables"),
             ("3", "Skills", "Set the expertise"), ("4", "Details", "Agree time + budget"),
             ("5", "Review", "Check the brief")]
    y = 3.42
    for number, label, detail in items:
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.43), Inches(y), Inches(0.31), Inches(0.31))
        circle.fill.solid(); circle.fill.fore_color.rgb = RED_LIGHT; circle.line.color.rgb = ORANGE
        add_text(slide, 9.43, y + 0.04, 0.31, 0.17, number, 8.5, ORANGE, True, PP_ALIGN.CENTER)
        add_text(slide, 9.9, y - 0.01, 1.6, 0.2, label, 11, NAVY, True)
        add_text(slide, 9.9, y + 0.2, 2.2, 0.2, detail, 9, MUTED)
        y += 0.62
    screen_box(slide, wizard, 496, 273, 1090, 817)
    screen_box(slide, wizard, 1143, 254, 1400, 1098)
    screen_box(slide, wizard, 972, 1015, 1090, 1073)
    callout(slide, 9.42, 6.08, 2.7, "Next step keeps momentum", 3)
    arrow(slide, 9.42, 6.27, *screen_point(wizard, 1030, 1045), RED, 1.7)
    add_footer(slide, 2)
    return slide


def slide_compare(prs):
    slide = add_slide_base(prs)
    add_header(slide, 3, "Compare evidence, then open a profile", "Matches are presented side by side so the choice is explainable.")
    matches = add_screen(slide, SCREENS["matches"], 0.62, 1.62, 6.12, 5.15)
    profile = add_screen(slide, SCREENS["profile"], 7.2, 1.62, 5.52, 5.15)
    screen_box(slide, matches, 295, 1120, 620, 1205)
    screen_box(slide, profile, 1100, 790, 1390, 900)
    callout(slide, 0.88, 6.84, 2.55, "Open the strongest match", 1)
    arrow(slide, 3.48, 7.03, *screen_point(matches, 460, 1165), RED, 1.7)
    callout(slide, 8.0, 6.84, 2.5, "Review fit + evidence", 2)
    arrow(slide, 10.55, 7.03, *screen_point(profile, 1245, 845), RED, 1.7)
    arrow(slide, 6.76, 4.12, 7.14, 4.12, RED, 2.4)
    add_footer(slide, 3)
    return slide


def slide_confirm(prs):
    slide = add_slide_base(prs)
    add_header(slide, 4, "Confirm a shared project plan", "The confirmation screen makes scope, budget, timing and milestones visible before commitment.")
    confirmation = add_screen(slide, SCREENS["confirmation"], 0.72, 1.56, 8.35, 5.35)
    rounded_box(slide, 9.48, 1.72, 3.08, 4.98, WHITE, LINE, True, 1)
    add_text(slide, 9.8, 2.08, 2.45, 0.3, "What to check", 15, NAVY, True)
    checks = ["Deliverables match the brief", "Milestones set expectations", "Budget and deadline are clear", "The guarantee is labelled as a concept"]
    y = 2.7
    for i, item in enumerate(checks, start=1):
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.82), Inches(y), Inches(0.3), Inches(0.3))
        badge.fill.solid(); badge.fill.fore_color.rgb = RED_LIGHT; badge.line.color.rgb = ORANGE
        add_text(slide, 9.82, y + 0.04, 0.3, 0.17, str(i), 8, ORANGE, True, PP_ALIGN.CENTER)
        add_text(slide, 10.28, y - 0.01, 1.9, 0.48, item, 10.5, MUTED)
        y += 0.72
    screen_box(slide, confirmation, 1085, 1140, 1400, 1235)
    callout(slide, 9.8, 5.98, 2.35, "Confirm to open workspace", 5)
    arrow(slide, 9.8, 6.18, *screen_point(confirmation, 1240, 1185), RED, 1.7)
    add_footer(slide, 4)
    return slide


def slide_work(prs):
    slide = add_slide_base(prs)
    add_header(slide, 5, "Work in one shared workspace", "Progress, reviews, deliverables and communication stay together after the match.")
    workspace = add_screen(slide, SCREENS["workspace"], 0.62, 1.55, 8.58, 5.35)
    rounded_box(slide, 9.52, 1.74, 3.03, 4.9, WHITE, LINE, True, 1)
    add_text(slide, 9.84, 2.06, 2.35, 0.28, "Use the workspace to…", 15, NAVY, True)
    steps = [("Review", "Check the current milestone"), ("Approve", "Move progress forward"),
             ("Message", "Clarify scope together"), ("Deliver", "Open sample outputs")]
    y = 2.72
    for i, (label, detail) in enumerate(steps, start=1):
        add_text(slide, 9.85, y, 0.32, 0.22, f"0{i}", 9, ORANGE, True)
        add_text(slide, 10.3, y - 0.01, 1.8, 0.22, label, 11, NAVY, True)
        add_text(slide, 10.3, y + 0.22, 1.85, 0.35, detail, 9.5, MUTED)
        y += 0.82
    screen_box(slide, workspace, 820, 500, 1085, 635)
    screen_box(slide, workspace, 1045, 530, 1330, 620)
    screen_box(slide, workspace, 1050, 650, 1320, 780)
    callout(slide, 9.82, 6.05, 2.38, "Every action has context", 6)
    arrow(slide, 9.82, 6.25, *screen_point(workspace, 1150, 570), RED, 1.7)
    add_footer(slide, 5)
    return slide


def slide_reflect(prs):
    slide = add_slide_base(prs)
    add_header(slide, 6, "Reflect and improve the next iteration", "Expert feedback closes the loop and keeps observations attached to the prototype experience.")
    feedback = add_screen(slide, SCREENS["feedback"], 0.62, 1.57, 8.55, 5.32)
    rounded_box(slide, 9.55, 1.72, 3.0, 4.93, WHITE, LINE, True, 1)
    add_text(slide, 9.87, 2.05, 2.35, 0.28, "A simple review loop", 15, NAVY, True)
    loop = [("Observe", "Rate clarity, evidence and next steps"), ("Capture", "Write one concrete improvement"), ("Download", "Share a readable Markdown note")]
    y = 2.82
    for i, (label, detail) in enumerate(loop, start=1):
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.88), Inches(y), Inches(0.31), Inches(0.31))
        badge.fill.solid(); badge.fill.fore_color.rgb = RED_LIGHT; badge.line.color.rgb = ORANGE
        add_text(slide, 9.88, y + 0.04, 0.31, 0.17, str(i), 8, ORANGE, True, PP_ALIGN.CENTER)
        add_text(slide, 10.35, y - 0.01, 1.75, 0.22, label, 11, NAVY, True)
        add_text(slide, 10.35, y + 0.22, 1.9, 0.45, detail, 9.5, MUTED)
        y += 0.9
    screen_box(slide, feedback, 300, 340, 1060, 400)
    screen_box(slide, feedback, 300, 900, 1070, 1037)
    callout(slide, 9.87, 6.05, 2.25, "Save feedback locally", 7)
    arrow(slide, 9.87, 6.25, *screen_point(feedback, 690, 970), RED, 1.7)
    add_footer(slide, 6)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "TaskJuvo prototype walkthrough"
    prs.core_properties.subject = "Annotated navigation guide for expert review"
    prs.core_properties.author = "TaskJuvo"
    add_cover(prs)
    slide_start(prs)
    slide_define(prs)
    slide_compare(prs)
    slide_confirm(prs)
    slide_work(prs)
    slide_reflect(prs)
    output = ARTIFACTS / "TaskJuvo-walkthrough.pptx"
    prs.save(output)
    print(f"Created {output} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
